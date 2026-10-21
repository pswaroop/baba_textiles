"""Builds the India Post bulk-booking Excel workbook (the same column layout
as India Post's own `ArticleDetails` / `PickupAddress` template — see
`reference.xlsx`) from a set of paid orders, for staff to upload via India
Post's web portal or hand off to their SFTP flow.

This is a separate channel from the JSON `process-articles` API integration
in services.py/operations.py — but it shares the same Shipment record as its
source of truth (via `prepare_shipment_for_order`), so a barcode reserved
here won't be re-issued to a different order if it's later booked through
the API instead, and vice versa.
"""
from datetime import timedelta

from django.conf import settings
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from .operations import ShipmentBookingError, prepare_shipment_for_order
from .services import IndiaPostClient, split_address_lines


def _digits_only(value, length=None):
    value = "".join(ch for ch in str(value or "") if ch.isdigit())
    return value[-length:] if length else value

ARTICLE_DETAILS_HEADERS = [
    "SERIAL NUMBER", "BARCODE NO", "PHYSICAL WEIGHT", "SHAPE OF ARTICLE",
    "LENGTH ", "BREADTH/DIAMETER", "HEIGHT", "PRIORITY FLAG",
    "DELIVERY INSTRUCTION", "INSTRUCTION RTS",
    "SENDER NAME", "SENDER COMPANY", "SENDER ADD LINE 1", "SENDER ADD LINE 2",
    "SENDER CITY", "SENDER STATE", "SENDER PINCODE", "SENDER EMAILID",
    "SENDER ALT CONTACT", "SENDER KYC", "SENDER TAX REFERENCE",
    "RECEIVER NAME", "RECEIVER COMPANY", "RECEIVER ADD LINE 1", "RECEIVER ADD LINE 2",
    "RECEIVER CITY", "RECEIVER STATE", "RECEIVER PINCODE", "RECEIVER EMAILID",
    "RECEIVER ALT CONTACT", "RECEIVER KYC", "RECEIVER TAX REFERENCE",
    "ALT ADDRESS FLAG", "PICKUP ADDRESS FLAG", "DROP OFF PINCODE",
    "DROPOFF/PICKUP OFFICE ID", "SENDER MOBILE NO", "RECEIVER MOBILE NO",
    "PREPAYMENT CODE", "VALUE OF PREPAYMENT", "CODR/COD", "VALUE FOR CODR/COD",
    "INSURANCE TYPE", "VALUE OF INSURANCE", "ACK", "REGISTRATION",
    "OTP BASED DELIVERY", "BULK REFERENCE",
]

PICKUP_ADDRESS_HEADERS = [
    "serial_no", "addressee_name", "company_name", "address_line1", "address_line2",
    "address_line3", "city", "state", "pincode", "email_id", "alt_contact_no",
    "mobile_no", "pickup_schedule_slot", "pickup_schedule_date",
]


def _next_pickup_date():
    return (timezone.localtime() + timedelta(days=1)).date()


def _order_row(serial_no, order, shipment):
    is_pickup = settings.INDIAPOST_PICKUP_OR_DROPOFF == "PICKUP"
    address_line1, address_line2 = split_address_lines(order.address)
    receiver_email = order.user.email if order.user_id and order.user.email else ""

    return [
        serial_no,
        shipment.barcode_no,
        shipment.physical_weight_grams,
        IndiaPostClient().shape_of_article_for(shipment.service_type),
        shipment.length_cm,
        shipment.breadth_cm,
        shipment.height_cm,
        False,  # PRIORITY FLAG
        "ND",   # DELIVERY INSTRUCTION — Normal Delivery
        "",     # INSTRUCTION RTS
        settings.INDIAPOST_SENDER_NAME,
        settings.INDIAPOST_SENDER_COMPANY,
        settings.INDIAPOST_SENDER_ADDRESS_LINE1,
        settings.INDIAPOST_SENDER_ADDRESS_LINE2,
        settings.INDIAPOST_SENDER_CITY,
        settings.INDIAPOST_SENDER_STATE,
        settings.INDIAPOST_SENDER_PINCODE,
        "",     # SENDER EMAILID
        settings.INDIAPOST_SENDER_MOBILE,
        "",     # SENDER KYC
        "",     # SENDER TAX REFERENCE
        order.customer_name,
        "",     # RECEIVER COMPANY
        address_line1,
        address_line2,
        order.city or "",
        order.state or "",
        _digits_only(order.pincode, 6),
        receiver_email,
        "",     # RECEIVER ALT CONTACT
        "",     # RECEIVER KYC
        "",     # RECEIVER TAX REFERENCE
        False,  # ALT ADDRESS FLAG
        is_pickup,
        "" if is_pickup else _digits_only(order.pincode, 6),  # DROP OFF PINCODE
        "",     # DROPOFF/PICKUP OFFICE ID — left blank; resolved automatically
                # by India Post from the pincode, same as our real booking history.
        settings.INDIAPOST_SENDER_MOBILE,
        _digits_only(order.phone, 10),
        "",     # PREPAYMENT CODE
        "",     # VALUE OF PREPAYMENT
        "",     # CODR/COD
        "",     # VALUE FOR CODR/COD
        "",     # INSURANCE TYPE
        "",     # VALUE OF INSURANCE
        False,  # ACK
        False,  # REGISTRATION
        False,  # OTP BASED DELIVERY
        order.order_number,
    ]


def _style_header(ws, headers):
    bold_red = Font(bold=True, color="FFFF0000")
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = bold_red
        ws.column_dimensions[get_column_letter(col_idx)].width = max(14, len(header) + 2)


def build_shipment_export_workbook(orders):
    """Builds the export workbook for `orders`.

    Returns (workbook, exported_orders, skipped) where `skipped` is a list of
    (order_number, reason) for orders that couldn't be prepared (not paid,
    already booked/shipped elsewhere).
    """
    wb = Workbook()
    articles_ws = wb.active
    articles_ws.title = "ArticleDetails"
    _style_header(articles_ws, ARTICLE_DETAILS_HEADERS)

    exported_orders = []
    skipped = []
    serial_no = 1

    for order in orders:
        try:
            shipment = prepare_shipment_for_order(order)
        except ShipmentBookingError as exc:
            skipped.append((order.order_number, str(exc)))
            continue

        articles_ws.append(_order_row(serial_no, order, shipment))
        exported_orders.append(order)
        serial_no += 1

    if settings.INDIAPOST_PICKUP_OR_DROPOFF == "PICKUP" and exported_orders:
        pickup_ws = wb.create_sheet("PickupAddress")
        _style_header(pickup_ws, PICKUP_ADDRESS_HEADERS)
        pickup_ws.append([
            1,
            settings.INDIAPOST_SENDER_NAME,
            settings.INDIAPOST_SENDER_COMPANY,
            settings.INDIAPOST_SENDER_ADDRESS_LINE1,
            settings.INDIAPOST_SENDER_ADDRESS_LINE2,
            "",
            settings.INDIAPOST_SENDER_CITY,
            settings.INDIAPOST_SENDER_STATE,
            settings.INDIAPOST_SENDER_PINCODE,
            "",
            settings.INDIAPOST_SENDER_MOBILE,
            settings.INDIAPOST_SENDER_MOBILE,
            settings.INDIAPOST_PICKUP_SLOT,
            _next_pickup_date(),
        ])
        pickup_ws["N2"].number_format = "DD-MM-YYYY"

    return wb, exported_orders, skipped
